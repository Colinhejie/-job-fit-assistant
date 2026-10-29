# -*- coding: utf-8 -*-
"""DOCX resume renderer: single column, Chinese-first, no forced shrinking.

The renderer never decides to squeeze the document into one page by dropping the
font below `min_body_pt`. The exporter above it may try progressively smaller
sizes, but stops at the floor and reports the resulting page count instead.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from ..errors import DependencyMissing, ExportError

DEFAULT_STYLE: dict[str, Any] = {
    "name": "default-zh-single-column",
    "page": {"size": "A4", "margins_cm": {"top": 1.5, "bottom": 1.5, "left": 1.7, "right": 1.7}},
    "fonts": {"latin": "Calibri", "east_asia": "Microsoft YaHei"},
    "sizes": {
        "name": 19, "target_title": 10.5, "contact": 9.5,
        "section": 10.5, "body": 10, "meta": 9.5,
    },
    "min_body_pt": 9.5,
    "colors": {"text": "1A1A1A", "accent": "1F4E5F", "muted": "4A4A4A"},
    "layout": {
        "sections": ["summary", "skills", "experience", "education", "achievements"],
        "bullet_char": "·",
        "section_rule": True,
        "photo": {"enabled": False, "width_cm": 2.6},
    },
    "labels": {
        "summary": "个人简介", "skills": "核心技能", "experience": "工作经历",
        "education": "教育背景", "achievements": "证书与荣誉",
    },
    "rules": [
        "正文最小字号 9.5pt，不通过缩小字号强行压成一页",
        "单栏、无表格化排版，便于 ATS 与人工阅读",
        "公司/职位/时间来自事实档案，改写不得修改",
    ],
}


def load_style(style_path: Path | None) -> dict:
    if style_path is None or not Path(style_path).exists():
        return json.loads(json.dumps(DEFAULT_STYLE))
    data = json.loads(Path(style_path).read_text(encoding="utf-8"))
    merged = json.loads(json.dumps(DEFAULT_STYLE))

    def merge(base: dict, override: dict) -> None:
        for key, value in override.items():
            if isinstance(value, dict) and isinstance(base.get(key), dict):
                merge(base[key], value)
            else:
                base[key] = value

    merge(merged, data)
    return merged


def render(
    doc_model: dict,
    style: dict,
    out_path: Path,
    *,
    photo_path: Path | None = None,
    body_pt: float | None = None,
) -> dict:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as exc:  # pragma: no cover
        raise DependencyMissing(
            "未安装 python-docx，无法生成 DOCX。",
            hint="pip install -r scripts/requirements.txt",
        ) from exc

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    fonts = style.get("fonts", {})
    latin = fonts.get("latin", "Calibri")
    east_asia = fonts.get("east_asia", "Microsoft YaHei")
    sizes = dict(style.get("sizes", {}))
    if body_pt:
        delta = float(body_pt) - float(sizes.get("body", 10))
        for key in ("body", "meta", "contact"):
            sizes[key] = round(float(sizes.get(key, 10)) + delta, 2)
    colors = style.get("colors", {})
    layout = style.get("layout", {})
    labels = style.get("labels", {})
    min_body = float(style.get("min_body_pt", 9.5))

    if float(sizes.get("body", 10)) < min_body - 1e-6:
        raise ExportError(
            f"请求的正文字号 {sizes['body']}pt 低于模板下限 {min_body}pt。",
            hint="不允许通过过小字号强行压缩页数；请压缩内容或增加页数上限。",
        )

    document = Document()
    _setup_page(document, style, Cm)
    _setup_normal(document, latin, east_asia, sizes, colors, Pt, RGBColor, qn)

    usable_width_cm = _usable_width_cm(style)

    # ---------------------------------------------------------------- header
    name = str(doc_model.get("name") or "").strip()
    if name:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(name)
        _font(run, latin, east_asia, sizes.get("name", 19), True, colors.get("text"), Pt, qn)

    target_title = str(doc_model.get("target_title") or "").strip()
    if target_title:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(target_title)
        _font(run, latin, east_asia, sizes.get("target_title", 10.5), False,
              colors.get("accent"), Pt, qn)

    contact_bits = [c["text"] for c in doc_model.get("contacts") or [] if c.get("text")]
    if contact_bits:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run("  |  ".join(contact_bits))
        _font(run, latin, east_asia, sizes.get("contact", 9.5), False, colors.get("muted"), Pt, qn)

    photo_cfg = layout.get("photo", {}) or {}
    if photo_cfg.get("enabled") and photo_path and Path(photo_path).exists():
        try:
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.add_run().add_picture(str(photo_path), width=Cm(float(photo_cfg.get("width_cm", 2.6))))
        except Exception:
            # a missing/corrupt photo must not abort the resume export
            pass

    # -------------------------------------------------------------- sections
    order = layout.get("sections") or ["summary", "skills", "experience", "education", "achievements"]

    def section(title: str) -> None:
        p = document.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title)
        _font(run, latin, east_asia, sizes.get("section", 10.5), True,
              colors.get("accent"), Pt, qn)
        if layout.get("section_rule", True):
            _bottom_border(p, colors.get("accent", "1F4E5F"), OxmlElement, qn)

    if "summary" in order and doc_model.get("summary"):
        section(labels.get("summary", "个人简介"))
        _body(document, doc_model["summary"]["text"], latin, east_asia, sizes, colors, Pt, qn)

    if "skills" in order and doc_model.get("skills"):
        section(labels.get("skills", "核心技能"))
        for skill in doc_model["skills"]:
            _bullet(document, skill["text"], latin, east_asia, sizes, colors, Pt, qn,
                    bullet_char=layout.get("bullet_char", "·"))

    if "experience" in order and doc_model.get("experiences"):
        section(labels.get("experience", "工作经历"))
        for exp in doc_model["experiences"]:
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.tab_stops.add_tab_stop(Cm(usable_width_cm), WD_TAB_ALIGNMENT.RIGHT)
            head_parts = [exp.get("company", ""), exp.get("title", "")]
            head = "  |  ".join(x for x in head_parts if x)
            run = p.add_run(head)
            _font(run, latin, east_asia, sizes.get("body", 10), True, colors.get("text"), Pt, qn)
            period = exp.get("period") or ""
            if period:
                run = p.add_run("\t" + period)
                _font(run, latin, east_asia, sizes.get("meta", 9.5), False,
                      colors.get("muted"), Pt, qn)
            for bullet in exp.get("bullets") or []:
                text = bullet["text"]
                if bullet.get("verification_status") == "pending_confirmation":
                    text = f"【待确认】{text}"
                _bullet(document, text, latin, east_asia, sizes, colors, Pt, qn,
                        bullet_char=layout.get("bullet_char", "·"))

    if "education" in order and doc_model.get("education"):
        section(labels.get("education", "教育背景"))
        for edu in doc_model["education"]:
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(usable_width_cm), WD_TAB_ALIGNMENT.RIGHT)
            head = "  |  ".join(x for x in (edu.get("school"), edu.get("degree"),
                                            edu.get("major")) if x) or edu.get("text", "")
            run = p.add_run(head)
            _font(run, latin, east_asia, sizes.get("body", 10), False, colors.get("text"), Pt, qn)
            if edu.get("date_raw") and edu.get("school"):
                run = p.add_run("\t" + edu["date_raw"])
                _font(run, latin, east_asia, sizes.get("meta", 9.5), False,
                      colors.get("muted"), Pt, qn)

    if "achievements" in order and doc_model.get("achievements"):
        section(labels.get("achievements", "证书与荣誉"))
        for item in doc_model["achievements"]:
            _bullet(document, item["text"], latin, east_asia, sizes, colors, Pt, qn,
                    bullet_char=layout.get("bullet_char", "·"))

    document.save(str(out_path))
    return {
        "path": str(out_path),
        "body_pt": sizes.get("body"),
        "paragraph_count": len(document.paragraphs),
        "size_bytes": out_path.stat().st_size,
    }


# ---------------------------------------------------------------------------
# primitives
# ---------------------------------------------------------------------------


def _setup_page(document, style: dict, Cm) -> None:
    page = style.get("page", {})
    margins = page.get("margins_cm", {})
    for section in document.sections:
        size = (page.get("size") or "A4").upper()
        if size == "A4":
            section.page_width = Cm(21.0)
            section.page_height = Cm(29.7)
        elif size == "LETTER":
            section.page_width = Cm(21.59)
            section.page_height = Cm(27.94)
        section.top_margin = Cm(float(margins.get("top", 1.5)))
        section.bottom_margin = Cm(float(margins.get("bottom", 1.5)))
        section.left_margin = Cm(float(margins.get("left", 1.7)))
        section.right_margin = Cm(float(margins.get("right", 1.7)))


def _usable_width_cm(style: dict) -> float:
    page = style.get("page", {})
    margins = page.get("margins_cm", {})
    width = 21.0 if (page.get("size") or "A4").upper() == "A4" else 21.59
    return round(width - float(margins.get("left", 1.7)) - float(margins.get("right", 1.7)), 2)


def _setup_normal(document, latin, east_asia, sizes, colors, Pt, RGBColor, qn) -> None:
    style = document.styles["Normal"]
    style.font.name = latin
    style.font.size = Pt(float(sizes.get("body", 10)))
    style.font.color.rgb = RGBColor.from_string(colors.get("text", "1A1A1A").lstrip("#"))
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.12


def _font(run, latin, east_asia, size_pt, bold, color, Pt, qn) -> None:
    run.font.size = Pt(float(size_pt))
    run.font.bold = bool(bold)
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    if color:
        from docx.shared import RGBColor

        run.font.color.rgb = RGBColor.from_string(str(color).lstrip("#"))


def _body(document, text, latin, east_asia, sizes, colors, Pt, qn) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(str(text))
    _font(run, latin, east_asia, sizes.get("body", 10), False, colors.get("text"), Pt, qn)


def _bullet(document, text, latin, east_asia, sizes, colors, Pt, qn, *, bullet_char="·") -> None:
    from docx.shared import Cm

    p = document.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.42)
    pf.first_line_indent = Cm(-0.42)
    pf.space_after = Pt(1)
    run = p.add_run(f"{bullet_char} ")
    _font(run, latin, east_asia, sizes.get("body", 10), False, colors.get("muted"), Pt, qn)
    run = p.add_run(str(text))
    _font(run, latin, east_asia, sizes.get("body", 10), False, colors.get("text"), Pt, qn)


def _bottom_border(paragraph, color: str, OxmlElement, qn) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), str(color).lstrip("#"))
    pbdr.append(bottom)
    ppr.append(pbdr)
